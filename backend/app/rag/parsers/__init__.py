"""Document parsers: turn uploaded files into plain text with a stable line structure."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import PurePosixPath

from app.core.exceptions import ErrorCode, ValidationFailed
from app.projects.scanner import CODE_LANGUAGES, detect_language

SUPPORTED_DOCUMENT_TYPES = {"markdown", "text", "pdf", "docx", "code", "config", "schema"}
_CONFIG_LANGS = {"json", "yaml", "toml", "ini", "xml", "dockerfile", "makefile", "dotenv", "requirements", "gomod"}


@dataclass
class ParsedDocument:
    text: str
    doc_type: str
    language: str | None = None
    # (1-based line number -> page number) for paginated formats
    page_starts: list[tuple[int, int]] = field(default_factory=list)

    def page_for_line(self, line: int) -> int | None:
        page = None
        for start_line, page_no in self.page_starts:
            if line >= start_line:
                page = page_no
        return page


def classify(path: str) -> tuple[str, str | None]:
    """Return (document_type, language) for a file path."""
    suffix = PurePosixPath(path).suffix.lower()
    if suffix == ".pdf":
        return "pdf", None
    if suffix == ".docx":
        return "docx", None
    language = detect_language(path)
    if language == "markdown" or suffix in {".md", ".mdx", ".rst"}:
        return "markdown", language or "markdown"
    if language == "sql":
        return "schema", "sql"
    if language in CODE_LANGUAGES:
        return "code", language
    if language in _CONFIG_LANGS:
        return "config", language
    return "text", language


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    return "\n".join(line.rstrip() for line in text.split("\n"))


def parse_pdf(data: bytes) -> ParsedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
    except (PdfReadError, ValueError, OSError) as exc:
        raise ValidationFailed("The PDF could not be read (corrupted or encrypted).", code=ErrorCode.VALIDATION_ERROR,
                               detail=str(exc)) from exc
    lines: list[str] = []
    page_starts: list[tuple[int, int]] = []
    for number, page in enumerate(reader.pages, start=1):
        page_starts.append((len(lines) + 1, number))
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # pypdf raises assorted errors for odd pages
            text = f"[page {number} could not be extracted: {type(exc).__name__}]"
        lines.extend(_clean(text).split("\n"))
    doc = ParsedDocument(text="\n".join(lines), doc_type="pdf", page_starts=page_starts)
    if not doc.text.strip():
        raise ValidationFailed("No text could be extracted from this PDF (it may be scanned images; OCR is not enabled).")
    return doc


def parse_docx(data: bytes) -> ParsedDocument:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, ValueError, KeyError) as exc:
        raise ValidationFailed("The DOCX file could not be read.", detail=str(exc)) from exc
    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        style = (para.style.name or "").lower() if para.style is not None else ""
        if text and style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            lines.append("#" * min(int(level), 6) + " " + text)
        else:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return ParsedDocument(text=_clean("\n".join(lines)), doc_type="docx", language="markdown")


def parse_bytes(path: str, data: bytes) -> ParsedDocument:
    doc_type, language = classify(path)
    if doc_type == "pdf":
        return parse_pdf(data)
    if doc_type == "docx":
        return parse_docx(data)
    if b"\x00" in data[:8192]:
        raise ValidationFailed(f"'{path}' looks like a binary file and cannot be indexed as text.")
    return ParsedDocument(text=_clean(data.decode("utf-8", errors="replace")), doc_type=doc_type, language=language)
